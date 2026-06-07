"""
ГосДок — Views подписей (apps/signatures/views.py)
Раздел 4.7 ТЗ

При подписи:
- Проверяем роль signer/owner/editor
- Сохраняем метаданные (время, IP, сертификат)
- Если все подписанты подписали → статус signed, блокируем редактирование
- Уведомляем всех участников (раздел 2.8 ТЗ)
"""

import logging

from django.shortcuts import get_object_or_404
from rest_framework import generics, permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.documents.models import Document
from apps.workspaces.models import WorkspaceMember
from .models import Signature
from .serializers import SignDocumentSerializer, SignatureSerializer

logger = logging.getLogger(__name__)


class SignDocumentView(APIView):
    """
    POST /api/v1/documents/{id}/sign/
    JWT + Signer — подписать документ.
    Раздел 2.7 ТЗ: подписанный документ блокируется от редактирования.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, pk):
        document = get_object_or_404(
            Document.objects.filter(
                workspace__members__user=request.user
            ).select_related("workspace"),
            pk=pk,
        )

        # Только подписанты и выше (раздел 2.2 ТЗ)
        member = WorkspaceMember.objects.filter(
            workspace=document.workspace,
            user=request.user,
            role__in=[
                WorkspaceMember.Role.OWNER,
                WorkspaceMember.Role.EDITOR,
                WorkspaceMember.Role.SIGNER,
            ],
        ).first()

        if not member:
            return Response(
                {"detail": "Только подписанты (signer/editor/owner) могут подписывать документы."},
                status=status.HTTP_403_FORBIDDEN,
            )

        # Подписанный документ заблокирован (раздел 2.7 ТЗ)
        if document.status == Document.DocumentStatus.SIGNED:
            return Response(
                {"detail": "Документ уже подписан и заблокирован от изменений."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Архивный документ нельзя подписывать
        if document.status == Document.DocumentStatus.ARCHIVED:
            return Response(
                {"detail": "Архивный документ нельзя подписать."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Раздел 2.3 ТЗ: подпись возможна только после завершения всех подзадач.
        pending_subtasks = document.subtasks.exclude(status="done").order_by("created_at")
        if pending_subtasks.exists():
            pending_titles = list(pending_subtasks.values_list("title", flat=True)[:5])
            return Response(
                {
                    "detail": (
                        "Нельзя подписать документ: сначала должны быть выполнены "
                        f"все подзадачи. Осталось незакрытых: {pending_subtasks.count()}."
                    ),
                    "pending_subtasks": pending_titles,
                },
                status=status.HTTP_409_CONFLICT,
            )

        # Проверяем, не подписал ли пользователь уже этот документ
        if Signature.objects.filter(document=document, user=request.user, is_valid=True).exists():
            return Response(
                {"detail": "Вы уже подписали этот документ."},
                status=status.HTTP_409_CONFLICT,
            )

        serializer = SignDocumentSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        # Получаем IP-адрес (метаданные подписи, раздел 3.7 ТЗ)
        ip_address = self._get_client_ip(request)

        signature = Signature.objects.create(
            document=document,
            user=request.user,
            signature_data=serializer.validated_data["signature_data"],
            certificate_id=serializer.validated_data.get("certificate_id", "") or "",
            ip_address=ip_address,
        )

        logger.info(
            "Документ '%s' подписан пользователем %s (IP: %s)",
            document.title, request.user.email, ip_address,
        )

        # Завершаем workflow-задачу подписанта: это двигает прогресс согласования
        # и через сигнал создаёт блок в блокчейне. Затем активируем следующий шаг.
        from django.utils import timezone
        from apps.tasks.models import Task as WfTask
        from apps.tasks.workflow import activate_next_task
        my_task = WfTask.objects.filter(
            document=document,
            assigned_to=request.user,
            status=WfTask.TaskStatus.IN_PROGRESS,
        ).order_by("step_order").first()
        if my_task:
            my_task.status = WfTask.TaskStatus.DONE
            my_task.completed_at = timezone.now()
            my_task.save(update_fields=["status", "completed_at"])
            activate_next_task(my_task)

        # Проверяем: все ли обязательные подписанты подписали
        all_signed = self._check_all_signed(document)
        if all_signed:
            document.status = Document.DocumentStatus.SIGNED
            document.save(update_fields=["status", "updated_at"])
            logger.info("Документ '%s' полностью подписан — статус SIGNED", document.title)

            # Уведомляем всех участников (раздел 2.8 ТЗ)
            from apps.tasks.workflow import notify_document_signed
            notify_document_signed(document, request.user)

        return Response(
            {
                "signature": SignatureSerializer(signature).data,
                "document_fully_signed": all_signed,
            },
            status=status.HTTP_201_CREATED,
        )

    @staticmethod
    def _get_client_ip(request) -> str:
        """Извлекает реальный IP с учётом reverse proxy."""
        x_forwarded_for = request.META.get("HTTP_X_FORWARDED_FOR")
        if x_forwarded_for:
            return x_forwarded_for.split(",")[0].strip()
        return request.META.get("REMOTE_ADDR", "0.0.0.0")

    @staticmethod
    def _check_all_signed(document) -> bool:
        """
        Проверяет, подписали ли документ все участники с ролью signer/owner.
        Возвращает True, если хотя бы один подписант есть и все подписали.
        """
        required_signers = WorkspaceMember.objects.filter(
            workspace=document.workspace,
            role__in=[WorkspaceMember.Role.OWNER, WorkspaceMember.Role.SIGNER],
        ).values_list("user_id", flat=True)

        if not required_signers:
            return False

        signed_users = Signature.objects.filter(
            document=document,
            is_valid=True,
        ).values_list("user_id", flat=True)

        return set(required_signers).issubset(set(signed_users))


class SignatureListView(generics.ListAPIView):
    """
    GET /api/v1/documents/{id}/signatures/
    JWT + Member — список подписей документа.
    """
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = SignatureSerializer
    pagination_class = None  # возвращаем массив напрямую

    def get_queryset(self):
        document = get_object_or_404(
            Document.objects.filter(workspace__members__user=self.request.user),
            pk=self.kwargs["pk"],
        )
        return document.signatures.select_related("user").order_by("signed_at")


class SignatureVerifyView(APIView):
    """
    GET /api/v1/signatures/{id}/verify/
    JWT — верификация подписи.
    Раздел 2.7 ТЗ.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        signature = get_object_or_404(
            Signature.objects.filter(
                document__workspace__members__user=request.user
            ).select_related("user", "document").distinct(),
            pk=pk,
        )

        return Response({
            "id": str(signature.id),
            "is_valid": signature.is_valid,
            "document": {
                "id": str(signature.document.id),
                "title": signature.document.title,
                "status": signature.document.status,
            },
            "signer": {
                "id": str(signature.user.id) if signature.user else None,
                "full_name": signature.user.full_name if signature.user else None,
                "email": signature.user.email if signature.user else None,
            },
            "signed_at": signature.signed_at,
            "ip_address": signature.ip_address,
            "certificate_id": signature.certificate_id or None,
        })


class SignPdfView(APIView):
    """
    POST /api/v1/documents/{id}/sign-pdf/
    JWT + Signer — впечатывает подпись в PDF на указанных координатах,
    сохраняет результат как новую версию, создаёт запись Signature и
    завершает workflow-шаг подписанта (+ блок в блокчейне).

    Body:
      signature_data: data URL PNG (по умолчанию — подпись из профиля пользователя)
      page:   индекс страницы (0-based)
      x, y, width, height: доли [0..1] от размера страницы (origin — верхний левый угол)
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, pk):
        import base64
        import io

        document = get_object_or_404(
            Document.objects.filter(workspace__members__user=request.user).select_related("workspace"),
            pk=pk,
        )

        if (document.file_type or "").lower() != "pdf":
            return Response({"detail": "Этот документ не PDF."}, status=status.HTTP_400_BAD_REQUEST)

        member = WorkspaceMember.objects.filter(
            workspace=document.workspace, user=request.user,
            role__in=[WorkspaceMember.Role.OWNER, WorkspaceMember.Role.EDITOR, WorkspaceMember.Role.SIGNER],
        ).first()
        if not member:
            return Response(
                {"detail": "Только подписанты (signer/editor/owner) могут подписывать документы."},
                status=status.HTTP_403_FORBIDDEN,
            )

        if document.status == Document.DocumentStatus.SIGNED:
            return Response({"detail": "Документ уже подписан и заблокирован."}, status=status.HTTP_400_BAD_REQUEST)
        if document.status == Document.DocumentStatus.ARCHIVED:
            return Response({"detail": "Архивный документ нельзя подписать."}, status=status.HTTP_400_BAD_REQUEST)

        pending = document.subtasks.exclude(status="done")
        if pending.exists():
            return Response(
                {"detail": f"Сначала должны быть выполнены все подзадачи (осталось {pending.count()})."},
                status=status.HTTP_409_CONFLICT,
            )

        if Signature.objects.filter(document=document, user=request.user, is_valid=True).exists():
            return Response({"detail": "Вы уже подписали этот документ."}, status=status.HTTP_409_CONFLICT)

        sig_data = request.data.get("signature_data") or getattr(request.user, "signature_data", "") or ""
        if not sig_data or "," not in sig_data:
            return Response(
                {"detail": "Нет данных подписи. Создайте свою подпись в профиле."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            page_index = int(request.data.get("page", 0))
            x = float(request.data.get("x", 0.1))
            y = float(request.data.get("y", 0.8))
            w = float(request.data.get("width", 0.25))
            h = float(request.data.get("height", 0.08))
        except (TypeError, ValueError):
            return Response({"detail": "Некорректные координаты подписи."}, status=status.HTTP_400_BAD_REQUEST)

        # 1. Скачиваем текущий PDF
        try:
            pdf_bytes = self._download_bytes(document.storage_key)
        except Exception as exc:
            logger.error("sign-pdf: download error doc=%s: %s", document.id, exc)
            return Response({"detail": "Не удалось загрузить PDF из хранилища."}, status=status.HTTP_502_BAD_GATEWAY)

        # 2. Впечатываем подпись через PyMuPDF
        try:
            import fitz  # PyMuPDF
            png = base64.b64decode(sig_data.split(",", 1)[1])
            pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
            page_index = max(0, min(page_index, pdf.page_count - 1))
            page = pdf[page_index]
            pr = page.rect
            rect = fitz.Rect(x * pr.width, y * pr.height, (x + w) * pr.width, (y + h) * pr.height)
            page.insert_image(rect, stream=png, keep_proportion=True, overlay=True)
            out_bytes = pdf.tobytes(deflate=True)
            pdf.close()
        except Exception as exc:
            logger.error("sign-pdf: stamp error doc=%s: %s", document.id, exc)
            return Response({"detail": "Не удалось наложить подпись на PDF."}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        # 3. Сохраняем как новую (подписанную) версию
        from apps.documents.models import DocumentVersion
        from apps.documents.storage import generate_storage_key, upload_to_s3
        new_key = generate_storage_key(str(document.workspace_id), f"{document.title}_signed.pdf")
        if not upload_to_s3(io.BytesIO(out_bytes), new_key, content_type="application/pdf"):
            return Response({"detail": "Не удалось сохранить подписанный PDF."}, status=status.HTTP_503_SERVICE_UNAVAILABLE)

        version_number = document.versions.count() + 1
        version = DocumentVersion.objects.create(
            document=document, version_number=version_number,
            storage_key=new_key, checksum="pending", created_by=request.user,
        )
        document.storage_key = new_key
        document.current_version = version
        document.save(update_fields=["storage_key", "current_version", "updated_at"])

        # 4. Запись о подписи
        ip = (
            request.META.get("HTTP_X_FORWARDED_FOR", "").split(",")[0].strip()
            or request.META.get("REMOTE_ADDR", "0.0.0.0")
        )
        signature = Signature.objects.create(
            document=document, user=request.user,
            signature_data=sig_data, certificate_id="", ip_address=ip,
        )
        logger.info("PDF '%s' подписан пользователем %s (стр. %s)", document.title, request.user.email, page_index)

        # 5. Завершаем workflow-задачу подписанта (прогресс + блок в блокчейне через сигнал)
        from django.utils import timezone
        from apps.tasks.models import Task as WfTask
        from apps.tasks.workflow import activate_next_task
        my_task = WfTask.objects.filter(
            document=document, assigned_to=request.user, status=WfTask.TaskStatus.IN_PROGRESS,
        ).order_by("step_order").first()
        if my_task:
            my_task.status = WfTask.TaskStatus.DONE
            my_task.completed_at = timezone.now()
            my_task.save(update_fields=["status", "completed_at"])
            activate_next_task(my_task)

        # 6. Все подписали?
        all_signed = self._check_all_signed(document)
        if all_signed:
            document.status = Document.DocumentStatus.SIGNED
            document.save(update_fields=["status", "updated_at"])
            from apps.tasks.workflow import notify_document_signed
            notify_document_signed(document, request.user)

        return Response(
            {
                "signature": SignatureSerializer(signature).data,
                "document_fully_signed": all_signed,
                "version_number": version_number,
            },
            status=status.HTTP_201_CREATED,
        )

    @staticmethod
    def _download_bytes(storage_key: str) -> bytes:
        from apps.documents.storage import _is_local_storage, _local_path, get_s3_client
        if _is_local_storage():
            with open(_local_path(storage_key), "rb") as f:
                return f.read()
        import io
        from django.conf import settings as s
        buf = io.BytesIO()
        get_s3_client().download_fileobj(s.AWS_STORAGE_BUCKET_NAME, storage_key, buf)
        return buf.getvalue()

    @staticmethod
    def _check_all_signed(document) -> bool:
        required = WorkspaceMember.objects.filter(
            workspace=document.workspace,
            role__in=[WorkspaceMember.Role.OWNER, WorkspaceMember.Role.SIGNER],
        ).values_list("user_id", flat=True)
        if not required:
            return False
        signed = Signature.objects.filter(document=document, is_valid=True).values_list("user_id", flat=True)
        return set(required).issubset(set(signed))


class SignDocxView(APIView):
    """
    POST /api/v1/documents/{id}/sign-docx/
    JWT + Signer — впечатывает подпись в .docx (блок подписи в конце документа
    через python-docx), сохраняет как новую версию, создаёт запись Signature
    и завершает workflow-шаг подписанта (+ блок в блокчейне).
    Документ остаётся Word — скачивается как .docx с подписью.

    Body: signature_data (data URL PNG, по умолчанию — подпись из профиля)
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, pk):
        import base64
        import io

        document = get_object_or_404(
            Document.objects.filter(workspace__members__user=request.user).select_related("workspace"),
            pk=pk,
        )

        if (document.file_type or "").lower() not in ("docx", "odt"):
            return Response({"detail": "Этот документ не Word (.docx)."}, status=status.HTTP_400_BAD_REQUEST)

        member = WorkspaceMember.objects.filter(
            workspace=document.workspace, user=request.user,
            role__in=[WorkspaceMember.Role.OWNER, WorkspaceMember.Role.EDITOR, WorkspaceMember.Role.SIGNER],
        ).first()
        if not member:
            return Response(
                {"detail": "Только подписанты (signer/editor/owner) могут подписывать документы."},
                status=status.HTTP_403_FORBIDDEN,
            )
        if document.status == Document.DocumentStatus.SIGNED:
            return Response({"detail": "Документ уже подписан и заблокирован."}, status=status.HTTP_400_BAD_REQUEST)
        if document.status == Document.DocumentStatus.ARCHIVED:
            return Response({"detail": "Архивный документ нельзя подписать."}, status=status.HTTP_400_BAD_REQUEST)

        pending = document.subtasks.exclude(status="done")
        if pending.exists():
            return Response(
                {"detail": f"Сначала должны быть выполнены все подзадачи (осталось {pending.count()})."},
                status=status.HTTP_409_CONFLICT,
            )
        if Signature.objects.filter(document=document, user=request.user, is_valid=True).exists():
            return Response({"detail": "Вы уже подписали этот документ."}, status=status.HTTP_409_CONFLICT)

        sig_data = request.data.get("signature_data") or getattr(request.user, "signature_data", "") or ""
        if not sig_data or "," not in sig_data:
            return Response(
                {"detail": "Нет данных подписи. Создайте свою подпись в профиле."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # 1. Скачиваем .docx
        try:
            docx_bytes = SignPdfView._download_bytes(document.storage_key)
        except Exception as exc:
            logger.error("sign-docx: download error doc=%s: %s", document.id, exc)
            return Response({"detail": "Не удалось загрузить файл."}, status=status.HTTP_502_BAD_GATEWAY)

        # 2. Впечатываем блок подписи в конец .docx
        try:
            from datetime import datetime
            from docx import Document as Docx
            from docx.shared import Inches, RGBColor

            png = base64.b64decode(sig_data.split(",", 1)[1])
            docx = Docx(io.BytesIO(docx_bytes))
            docx.add_paragraph()
            head = docx.add_paragraph()
            run = head.add_run("Электронная подпись")
            run.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            docx.add_picture(io.BytesIO(png), width=Inches(2.2))
            full_name = getattr(request.user, "full_name", "") or request.user.email
            docx.add_paragraph(f"Подписал(а): {full_name}    Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
            out = io.BytesIO()
            docx.save(out)
            out_bytes = out.getvalue()
        except Exception as exc:
            logger.error("sign-docx: embed error doc=%s: %s", document.id, exc)
            return Response({"detail": "Не удалось вставить подпись в документ."}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        # 3. Сохраняем как новую версию (.docx)
        from apps.documents.models import DocumentVersion
        from apps.documents.storage import generate_storage_key, get_content_type, upload_to_s3
        new_key = generate_storage_key(str(document.workspace_id), f"{document.title}_signed.docx")
        if not upload_to_s3(io.BytesIO(out_bytes), new_key, content_type=get_content_type("x.docx")):
            return Response({"detail": "Не удалось сохранить подписанный файл."}, status=status.HTTP_503_SERVICE_UNAVAILABLE)

        version_number = document.versions.count() + 1
        version = DocumentVersion.objects.create(
            document=document, version_number=version_number,
            storage_key=new_key, checksum="pending", created_by=request.user,
        )
        document.storage_key = new_key
        document.current_version = version
        document.save(update_fields=["storage_key", "current_version", "updated_at"])

        # 4. Запись подписи
        ip = (
            request.META.get("HTTP_X_FORWARDED_FOR", "").split(",")[0].strip()
            or request.META.get("REMOTE_ADDR", "0.0.0.0")
        )
        signature = Signature.objects.create(
            document=document, user=request.user,
            signature_data=sig_data, certificate_id="", ip_address=ip,
        )
        logger.info("DOCX '%s' подписан пользователем %s", document.title, request.user.email)

        # 5. Завершаем задачу подписанта (+ блокчейн через сигнал) и активируем следующий шаг
        from django.utils import timezone
        from apps.tasks.models import Task as WfTask
        from apps.tasks.workflow import activate_next_task
        my_task = WfTask.objects.filter(
            document=document, assigned_to=request.user, status=WfTask.TaskStatus.IN_PROGRESS,
        ).order_by("step_order").first()
        if my_task:
            my_task.status = WfTask.TaskStatus.DONE
            my_task.completed_at = timezone.now()
            my_task.save(update_fields=["status", "completed_at"])
            activate_next_task(my_task)

        # 6. Все подписали?
        all_signed = SignPdfView._check_all_signed(document)
        if all_signed:
            document.status = Document.DocumentStatus.SIGNED
            document.save(update_fields=["status", "updated_at"])
            from apps.tasks.workflow import notify_document_signed
            notify_document_signed(document, request.user)

        return Response(
            {
                "signature": SignatureSerializer(signature).data,
                "document_fully_signed": all_signed,
                "version_number": version_number,
            },
            status=status.HTTP_201_CREATED,
        )
